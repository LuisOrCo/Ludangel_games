from fastapi import APIRouter, Depends, HTTPException, status, Query, Response
from sqlalchemy.orm import Session
from sqlalchemy import func, DATE
from typing import Dict, Any, List, Optional
from io import BytesIO
import datetime

from ..database import get_db
from ..models import Venta, DetalleVenta, Producto, Usuario, PQR
from ..security import require_roles

router = APIRouter(prefix="/reportes", tags=["Reportes y Estadísticas"])


@router.get("/dashboard", response_model=Dict[str, Any])
def obtener_metricas_dashboard(
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles([1, 2]))
):
    total_usuarios = db.query(func.count(Usuario.id_usuario)).scalar() or 0
    total_clientes = db.query(func.count(Usuario.id_usuario)).filter(Usuario.id_rol == 3).scalar() or 0
    total_empleados = db.query(func.count(Usuario.id_usuario)).filter(Usuario.id_rol == 2).scalar() or 0

    total_productos = db.query(func.count(Producto.id_producto)).scalar() or 0
    productos_activos = db.query(func.count(Producto.id_producto)).filter(Producto.estado == True).scalar() or 0
    stock_critico = db.query(func.count(Producto.id_producto)).filter(Producto.stock <= 5, Producto.estado == True).scalar() or 0

    total_ventas = db.query(func.count(Venta.id_venta)).scalar() or 0
    ingresos_totales = db.query(func.sum(Venta.total)).filter(Venta.estado == "Completada").scalar() or 0.0

    pqrs_pendientes = db.query(func.count(PQR.id_pqr)).filter(PQR.estado == "Pendiente").scalar() or 0
    total_pqrs = db.query(func.count(PQR.id_pqr)).scalar() or 0

    return {
        "usuarios": {
            "total": total_usuarios,
            "clientes": total_clientes,
            "empleados": total_empleados
        },
        "productos": {
            "total": total_productos,
            "activos": productos_activos,
            "stock_critico": stock_critico
        },
        "ventas": {
            "total_ventas": total_ventas,
            "ingresos_totales": float(ingresos_totales)
        },
        "pqrs": {
            "total": total_pqrs,
            "pendientes": pqrs_pendientes
        }
    }


@router.get("/productos-mas-vendidos", response_model=List[Dict[str, Any]])
def obtener_productos_mas_vendidos(
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles([1, 2]))
):
    resultados = (
        db.query(
            Producto.id_producto,
            Producto.nombre,
            Producto.precio,
            func.sum(DetalleVenta.cantidad).label("unidades_vendidas"),
            func.sum(DetalleVenta.subtotal).label("total_recaudado")
        )
        .join(DetalleVenta, Producto.id_producto == DetalleVenta.id_producto)
        .group_by(Producto.id_producto, Producto.nombre, Producto.precio)
        .order_by(func.sum(DetalleVenta.cantidad).desc())
        .limit(5)
        .all()
    )

    return [
        {
            "id_producto": r.id_producto,
            "nombre": r.nombre,
            "precio": float(r.precio),
            "unidades_vendidas": int(r.unidades_vendidas or 0),
            "total_recaudado": float(r.total_recaudado or 0.0)
        }
        for r in resultados
    ]


# ====================================================================
# REQ-04: REPORTE DIARIO DE VENTAS
# ====================================================================
@router.get("/diario")
def obtener_reporte_diario(
    fecha: Optional[str] = Query(None, description="Fecha en formato YYYY-MM-DD"),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles([1, 2]))
):
    """
    REQ-04: Genera el reporte diario de ventas con la información registrada en una fecha determinada.
    """
    target_date = fecha if fecha else datetime.date.today().strftime("%Y-%m-%d")

    ventas = (
        db.query(Venta)
        .filter(func.date(Venta.fecha_venta) == target_date)
        .order_by(Venta.fecha_venta.asc())
        .all()
    )

    total_recaudado = sum(float(v.total) for v in ventas)

    lista_ventas = []
    for v in ventas:
        items_str = ", ".join(
            [f"{d.producto.nombre if d.producto else 'Prod'} (x{d.cantidad})" for d in v.detalles]
        )
        lista_ventas.append({
            "id_venta": v.id_venta,
            "fecha": v.fecha_venta.strftime("%Y-%m-%d %H:%M:%S") if v.fecha_venta else "",
            "cliente": f"{v.usuario.nombre} {v.usuario.apellido}" if v.usuario else "Cliente",
            "correo": v.usuario.correo if v.usuario else "",
            "documento": f"{v.usuario.tipo_documento} {v.usuario.numero_documento}" if v.usuario else "",
            "items": items_str,
            "metodo_pago": v.metodo_pago,
            "total": float(v.total),
            "estado": v.estado
        })

    return {
        "fecha_reporte": target_date,
        "total_ventas": len(lista_ventas),
        "total_recaudado": total_recaudado,
        "ventas": lista_ventas
    }


# ====================================================================
# REQ-05: EXPORTACIÓN DEL REPORTE DIARIO EN PDF
# ====================================================================
@router.get("/diario/pdf")
def exportar_reporte_diario_pdf(
    fecha: Optional[str] = Query(None, description="Fecha en formato YYYY-MM-DD"),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles([1, 2]))
):
    target_date = fecha if fecha else datetime.date.today().strftime("%Y-%m-%d")
    ventas = (
        db.query(Venta)
        .filter(func.date(Venta.fecha_venta) == target_date)
        .order_by(Venta.fecha_venta.asc())
        .all()
    )

    total_recaudado = sum(float(v.total) for v in ventas)
    fecha_generacion = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.pdfgen import canvas
        from reportlab.lib import colors

        buffer = BytesIO()
        p = canvas.Canvas(buffer, pagesize=landscape(letter))
        width, height = landscape(letter)

        # Encabezado
        p.setFont("Helvetica-Bold", 18)
        p.setFillColor(colors.HexColor("#06B6D4"))
        p.drawString(40, height - 40, "LUDANGEL GAMES — REPORTE DIARIO DE VENTAS")

        p.setFont("Helvetica", 9)
        p.setFillColor(colors.HexColor("#64748B"))
        p.drawString(40, height - 55, "Proyecto: SENA ADSO | ID Proyecto: 3406204")
        p.drawRightString(width - 40, height - 40, f"Fecha Reporte: {target_date}")
        p.drawRightString(width - 40, height - 55, f"Generado: {fecha_generacion}")

        p.setStrokeColor(colors.HexColor("#334155"))
        p.line(40, height - 65, width - 40, height - 65)

        # Resumen general
        p.setFillColor(colors.HexColor("#0F172A"))
        p.rect(40, height - 105, width - 80, 32, fill=True, stroke=False)
        p.setFillColor(colors.HexColor("#FFFFFF"))
        p.setFont("Helvetica-Bold", 10)
        p.drawString(50, height - 90, f"Total Ventas Registradas: {len(ventas)}")
        p.drawRightString(width - 50, height - 90, f"Ingresos Totales del Día: ${total_recaudado:,.2f} COP")

        # Encabezado de la Tabla
        y = height - 125
        p.setFillColor(colors.HexColor("#1E293B"))
        p.rect(40, y - 5, width - 80, 18, fill=True, stroke=False)

        p.setFillColor(colors.HexColor("#38BDF8"))
        p.setFont("Helvetica-Bold", 9)
        p.drawString(45, y, "N° Venta")
        p.drawString(105, y, "Hora")
        p.drawString(160, y, "Cliente")
        p.drawString(310, y, "Productos / Cantidad")
        p.drawString(550, y, "Pago")
        p.drawString(640, y, "Total (COP)")
        p.drawString(730, y, "Estado")

        y -= 20
        p.setFont("Helvetica", 8)
        p.setFillColor(colors.HexColor("#0F172A"))

        if not ventas:
            p.setFillColor(colors.HexColor("#64748B"))
            p.drawString(45, y, "No se registraron ventas en esta fecha.")
        else:
            for v in ventas:
                if y < 50:  # Nueva página si se llena
                    p.showPage()
                    y = height - 50
                    p.setFont("Helvetica", 8)

                hora = v.fecha_venta.strftime("%H:%M:%S") if v.fecha_venta else ""
                cliente_nombre = f"{v.usuario.nombre} {v.usuario.apellido}" if v.usuario else "Cliente General"
                items_str = ", ".join([f"{d.producto.nombre if d.producto else 'Prod'} (x{d.cantidad})" for d in v.detalles])
                if len(items_str) > 42:
                    items_str = items_str[:40] + "..."

                p.drawString(45, y, f"#{v.id_venta:06d}")
                p.drawString(105, y, hora)
                p.drawString(160, y, cliente_nombre[:24])
                p.drawString(310, y, items_str)
                p.drawString(550, y, v.metodo_pago[:12])
                p.drawString(640, y, f"${float(v.total):,.2f}")
                p.drawString(730, y, v.estado)

                p.setStrokeColor(colors.HexColor("#E2E8F0"))
                p.setLineWidth(0.5)
                p.line(40, y - 4, width - 40, y - 4)
                y -= 16

        # Pie de página
        p.setFont("Helvetica-Oblique", 8)
        p.setFillColor(colors.HexColor("#94A3B8"))
        p.drawCentredString(width / 2.0, 20, "Sistema de Gestión LudAngel Games — Documento Oficial de Reporte Diario")

        p.showPage()
        p.save()

        pdf_bytes = buffer.getvalue()
        buffer.close()

        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=Reporte_Diario_Ventas_{target_date}.pdf"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar PDF: {str(e)}")


# ====================================================================
# REQ-06: EXPORTACIÓN DEL REPORTE DIARIO EN EXCEL (.XLSX)
# ====================================================================
@router.get("/diario/excel")
def exportar_reporte_diario_excel(
    fecha: Optional[str] = Query(None, description="Fecha en formato YYYY-MM-DD"),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles([1, 2]))
):
    target_date = fecha if fecha else datetime.date.today().strftime("%Y-%m-%d")
    ventas = (
        db.query(Venta)
        .filter(func.date(Venta.fecha_venta) == target_date)
        .order_by(Venta.fecha_venta.asc())
        .all()
    )

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Ventas {target_date}"

        # Título
        ws.merge_cells("A1:I1")
        ws["A1"] = f"LUDANGEL GAMES — REPORTE DIARIO DE VENTAS ({target_date})"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="06B6D4", end_color="06B6D4", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        # Encabezados de Columna
        headers = [
            "ID Venta", "Fecha y Hora", "Cliente", "Documento",
            "Correo", "Productos / Cantidad", "Método de Pago", "Valor Total (COP)", "Estado"
        ]
        ws.append([]) # Fila 2 vacía
        ws.append(headers) # Fila 3

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="38BDF8")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Datos
        total_recaudado = 0.0
        row_num = 4

        for v in ventas:
            items_str = ", ".join([f"{d.producto.nombre if d.producto else 'Producto'} (x{d.cantidad})" for d in v.detalles])
            cliente_nombre = f"{v.usuario.nombre} {v.usuario.apellido}" if v.usuario else "Cliente General"
            doc_str = f"{v.usuario.tipo_documento} {v.usuario.numero_documento}" if v.usuario else "N/A"
            correo_str = v.usuario.correo if v.usuario else "N/A"
            fecha_str = v.fecha_venta.strftime("%Y-%m-%d %H:%M:%S") if v.fecha_venta else ""

            ws.append([
                f"#{v.id_venta:06d}",
                fecha_str,
                cliente_nombre,
                doc_str,
                correo_str,
                items_str,
                v.metodo_pago,
                float(v.total),
                v.estado
            ])
            total_recaudado += float(v.total)
            row_num += 1

        # Fila Total
        ws.append([])
        ws.append(["TOTALES", "", "", "", "", "", f"Ventas: {len(ventas)}", total_recaudado, ""])
        last_row = ws.max_row
        ws.cell(row=last_row, column=1).font = Font(bold=True)
        ws.cell(row=last_row, column=8).font = Font(bold=True)
        ws.cell(row=last_row, column=8).number_format = '"$"#,##0.00'

        # Ajuste de ancho de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()

        return Response(
            content=excel_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=Reporte_Diario_Ventas_{target_date}.xlsx"}
        )

    except ImportError:
        # Fallback a CSV si openpyxl no estuviera
        import csv
        buffer = BytesIO()
        writer = csv.writer(buffer)
        writer.writerow(["ID Venta", "Fecha y Hora", "Cliente", "Documento", "Correo", "Productos", "Metodo Pago", "Total COP", "Estado"])
        for v in ventas:
            items_str = ", ".join([f"{d.producto.nombre if d.producto else 'Producto'} (x{d.cantidad})" for d in v.detalles])
            writer.writerow([
                v.id_venta,
                v.fecha_venta,
                f"{v.usuario.nombre} {v.usuario.apellido}" if v.usuario else "",
                v.usuario.numero_documento if v.usuario else "",
                v.usuario.correo if v.usuario else "",
                items_str,
                v.metodo_pago,
                v.total,
                v.estado
            ])

        return Response(
            content=buffer.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=Reporte_Diario_Ventas_{target_date}.csv"}
        )


# ====================================================================
# REQ-11 & REQ-13: DASHBOARD ANALÍTICO CON GRÁFICOS Y FILTROS AVANZADOS
# ====================================================================
@router.get("/analitica")
def obtener_analitica_ventas(
    fecha_inicio: Optional[str] = Query(None, description="Fecha inicio YYYY-MM-DD"),
    fecha_fin: Optional[str] = Query(None, description="Fecha fin YYYY-MM-DD"),
    id_producto: Optional[int] = Query(None, description="ID del producto a filtrar"),
    id_cliente: Optional[int] = Query(None, description="ID del cliente a filtrar"),
    estado: Optional[str] = Query(None, description="Estado de la venta"),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles([1, 2]))
):
    """
    REQ-11: Dashboard de ventas con visualización analítica (gráfico de barras, lineal e indicadores en Cards).
    REQ-13: Filtros por fecha inicial, fecha final, producto, cliente y estado.
    """
    query = db.query(Venta)

    if fecha_inicio:
        query = query.filter(func.date(Venta.fecha_venta) >= fecha_inicio)
    if fecha_fin:
        query = query.filter(func.date(Venta.fecha_venta) <= fecha_fin)
    if id_cliente:
        query = query.filter(Venta.id_usuario == id_cliente)
    if estado:
        query = query.filter(Venta.estado == estado)
    if id_producto:
        query = query.join(DetalleVenta, Venta.id_venta == DetalleVenta.id_venta).filter(DetalleVenta.id_producto == id_producto)

    ventas = query.order_by(Venta.fecha_venta.asc()).all()

    # 1. Indicadores numéricos en Cards
    total_ventas = len(ventas)
    ingresos_totales = sum(float(v.total) for v in ventas)
    promedio_venta = (ingresos_totales / total_ventas) if total_ventas > 0 else 0.0

    unidades_vendidas = 0
    for v in ventas:
        for d in v.detalles:
            if not id_producto or d.id_producto == id_producto:
                unidades_vendidas += d.cantidad

    # 2. Serie para Gráfico Lineal (Tendencia por día)
    ventas_por_dia = {}
    for v in ventas:
        dia_str = v.fecha_venta.strftime("%Y-%m-%d") if v.fecha_venta else "Sin fecha"
        if dia_str not in ventas_por_dia:
            ventas_por_dia[dia_str] = {"etiqueta": dia_str, "ingresos": 0.0, "ventas": 0}
        ventas_por_dia[dia_str]["ingresos"] += float(v.total)
        ventas_por_dia[dia_str]["ventas"] += 1

    serie_lineal = list(ventas_por_dia.values())

    # 3. Serie para Gráfico de Barras (Desglose por Producto)
    productos_dict = {}
    for v in ventas:
        for d in v.detalles:
            if id_producto and d.id_producto != id_producto:
                continue
            p_nombre = d.producto.nombre if d.producto else f"Producto #{d.id_producto}"
            if p_nombre not in productos_dict:
                productos_dict[p_nombre] = {"etiqueta": p_nombre, "unidades": 0, "ingresos": 0.0}
            productos_dict[p_nombre]["unidades"] += d.cantidad
            productos_dict[p_nombre]["ingresos"] += float(d.subtotal)

    serie_barras = list(productos_dict.values())
    serie_barras.sort(key=lambda x: x["unidades"], reverse=True)

    # 4. Listas para poblar los selects del filtro en el Frontend
    productos_lista = [
        {"id_producto": p.id_producto, "nombre": p.nombre}
        for p in db.query(Producto).filter(Producto.estado == True).all()
    ]
    clientes_lista = [
        {"id_usuario": u.id_usuario, "nombre": f"{u.nombre} {u.apellido}"}
        for u in db.query(Usuario).filter(Usuario.id_rol == 3).all()
    ]

    return {
        "cards": {
            "total_ventas": total_ventas,
            "ingresos_totales": ingresos_totales,
            "promedio_venta": promedio_venta,
            "unidades_vendidas": unidades_vendidas
        },
        "grafico_lineal": serie_lineal,
        "grafico_barras": serie_barras[:10], # Top 10 para barras
        "opciones_filtro": {
            "productos": productos_lista,
            "clientes": clientes_lista
        }
    }

